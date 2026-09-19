from __future__ import annotations

from dataclasses import dataclass
from datetime import datetime
from io import BytesIO
from pathlib import Path
from typing import Dict, List, Tuple

import numpy as np
import pandas as pd
import plotly.express as px
import plotly.graph_objects as go
import streamlit as st
from openpyxl import load_workbook
from scipy.optimize import linprog

# ============================================================
# CONFIGURATION
# ============================================================
st.set_page_config(
    page_title="Aruna Production Intelligence DSS",
    page_icon="◈",
    layout="wide",
    initial_sidebar_state="expanded",
)

NAVY = "#0E3448"
NAVY_2 = "#153F55"
TEAL = "#1697A6"
TEAL_2 = "#51B7BC"
AMBER = "#EE9B43"
RED = "#C75450"
GREEN = "#4B8E69"
INK = "#173746"
MUTED = "#6E818C"
BG = "#F5F7F9"
CARD = "#FFFFFF"
BORDER = "#E2E8EC"
SOFT_BLUE = "#EAF4F6"
SOFT_AMBER = "#FFF5EA"
SOFT_GREEN = "#EDF7F0"

st.markdown(
    f"""
    <style>
    #MainMenu, footer, header {{visibility:hidden;}}
    .stApp {{background:{BG}; color:{INK};}}
    .block-container {{padding-top:1.1rem; padding-bottom:2.5rem; max-width:1540px;}}
    [data-testid="stSidebar"] {{background:#F8FAFB; border-right:1px solid {BORDER};}}
    [data-testid="stSidebar"] .block-container {{padding-top:1rem;}}
    div[data-testid="stMetric"] {{background:{CARD}; border:1px solid {BORDER}; border-radius:16px; padding:.75rem .9rem; box-shadow:0 2px 9px rgba(14,52,72,.035);}}
    div[data-testid="stMetric"] label {{color:{MUTED}!important; font-size:.78rem!important; text-transform:uppercase; letter-spacing:.04em;}}
    div[data-testid="stMetric"] [data-testid="stMetricValue"] {{color:{NAVY}!important; font-weight:780!important;}}
    div[data-testid="stMetric"] [data-testid="stMetricDelta"] {{font-size:.76rem!important;}}
    div[data-testid="stDataFrame"] {{border:1px solid {BORDER}; border-radius:14px; overflow:hidden;}}
    .hero {{
        background:linear-gradient(115deg,{NAVY} 0%, #164E62 68%, {TEAL} 145%);
        color:white; border-radius:20px; padding:1.15rem 1.35rem 1.05rem 1.35rem;
        box-shadow:0 8px 24px rgba(14,52,72,.13); margin-bottom:1rem;
    }}
    .hero-kicker {{font-size:.73rem; text-transform:uppercase; letter-spacing:.13em; opacity:.78; font-weight:700;}}
    .hero-title {{font-size:1.75rem; line-height:1.1; font-weight:820; letter-spacing:-.035em; margin:.2rem 0 .28rem;}}
    .hero-sub {{font-size:.92rem; opacity:.86; max-width:980px;}}
    .hero-meta {{display:flex; gap:.48rem; flex-wrap:wrap; margin-top:.72rem;}}
    .hero-pill {{font-size:.72rem; padding:.28rem .58rem; border-radius:999px; background:rgba(255,255,255,.11); border:1px solid rgba(255,255,255,.18);}}
    .section-title {{font-size:1.05rem; font-weight:790; color:{NAVY}; margin:.35rem 0 .6rem; letter-spacing:-.015em;}}
    .section-kicker {{font-size:.72rem; font-weight:720; color:{TEAL}; text-transform:uppercase; letter-spacing:.09em; margin-bottom:.12rem;}}
    .panel {{background:{CARD}; border:1px solid {BORDER}; border-radius:16px; padding:.95rem 1rem; box-shadow:0 2px 10px rgba(14,52,72,.035);}}
    .panel-title {{font-weight:760; color:{NAVY}; font-size:.95rem; margin-bottom:.3rem;}}
    .panel-sub {{font-size:.78rem; color:{MUTED}; margin-bottom:.65rem;}}
    .signal {{border:1px solid {BORDER}; border-left:5px solid {TEAL}; border-radius:14px; padding:.78rem .9rem; background:{CARD}; margin:.45rem 0;}}
    .signal.warn {{border-left-color:{AMBER}; background:{SOFT_AMBER};}}
    .signal.good {{border-left-color:{GREEN}; background:{SOFT_GREEN};}}
    .signal.bad {{border-left-color:{RED}; background:#FCEEEE;}}
    .signal-title {{font-weight:760; color:{NAVY}; font-size:.9rem;}}
    .signal-body {{color:#536D79; font-size:.81rem; line-height:1.42; margin-top:.18rem;}}
    .badge {{display:inline-block; border-radius:999px; padding:.22rem .52rem; font-size:.68rem; font-weight:720; margin-right:.25rem;}}
    .badge-source {{background:{SOFT_BLUE}; color:#126B77;}}
    .badge-derived {{background:#EEF0FA; color:#4E5B9B;}}
    .badge-assumption {{background:{SOFT_AMBER}; color:#A56522;}}
    .badge-live {{background:{SOFT_GREEN}; color:#35724F;}}
    .data-card {{background:{CARD}; border:1px solid {BORDER}; border-radius:15px; padding:.85rem .9rem; min-height:122px;}}
    .data-card-name {{font-size:.82rem; color:{MUTED}; font-weight:690;}}
    .data-card-value {{font-size:1.45rem; color:{NAVY}; font-weight:820; letter-spacing:-.025em; margin:.16rem 0;}}
    .data-card-note {{font-size:.72rem; color:#728590; line-height:1.35;}}
    .action-card {{background:{CARD}; border:1px solid {BORDER}; border-radius:15px; padding:.85rem .9rem; min-height:135px;}}
    .action-title {{color:{NAVY}; font-weight:780; font-size:.89rem;}}
    .action-body {{font-size:.78rem; color:#5F7580; line-height:1.45; margin-top:.28rem;}}
    .mini {{font-size:.72rem; color:{MUTED};}}
    .sidebar-brand {{font-size:1.05rem; font-weight:840; color:{NAVY}; letter-spacing:-.02em;}}
    .sidebar-sub {{font-size:.72rem; color:{MUTED}; margin-bottom:.65rem;}}
    .scenario-card {{background:white; border:1px solid {BORDER}; border-radius:13px; padding:.65rem .72rem; margin:.55rem 0 .8rem;}}
    .scenario-name {{font-size:.8rem; color:{NAVY}; font-weight:760;}}
    .scenario-note {{font-size:.68rem; color:{MUTED}; margin-top:.2rem;}}
    .lineage {{display:flex; align-items:center; gap:.42rem; flex-wrap:wrap; margin:.55rem 0;}}
    .lineage-box {{background:white; border:1px solid {BORDER}; border-radius:10px; padding:.46rem .62rem; font-size:.73rem; color:{NAVY}; font-weight:700;}}
    .lineage-arrow {{color:{TEAL}; font-weight:900;}}
    .stButton > button {{border-radius:12px; border:0; background:{NAVY}; color:white; font-weight:760; padding:.58rem .9rem; box-shadow:none;}}
    .stButton > button:hover {{background:#174B61; color:white; border:0;}}
    .stDownloadButton > button {{border-radius:12px; border:1px solid {NAVY}; background:white; color:{NAVY}; font-weight:740;}}
    [data-testid="stForm"] {{background:white; border:1px solid {BORDER}; border-radius:16px; padding:.8rem .9rem;}}
    </style>
    """,
    unsafe_allow_html=True,
)

SKUS = [
    "Saku 16 OZ", "Steak 4 OZ", "Steak 6 OZ", "Steak 8 OZ",
    "Steak 10 OZ", "Poke 1.5 CM", "Medallion 2-3 OZ", "Ground Meat",
]
FRESH_GROUPS = ["30 UP Grade B", "30 UP Grade C", "20 UP Grade C", "14 UP Grade C"]
RESOURCE_GROUPS = FRESH_GROUPS + ["WIP ST/SK carry-in"]
DEFAULT_MODEL = Path(__file__).with_name("Model_Aruna_FINAL_Aligned.xlsx")


# ============================================================
# HELPERS
# ============================================================
def _num(v, fallback=0.0):
    try:
        if v is None:
            return float(fallback)
        return float(v)
    except Exception:
        return float(fallback)


def fmt_kg(v):
    return f"{v/1000:,.2f} t" if abs(v) >= 1000 else f"{v:,.0f} kg"


def fmt_rp(v):
    if abs(v) >= 1e9:
        return f"Rp{v/1e9:,.2f} miliar"
    if abs(v) >= 1e6:
        return f"Rp{v/1e6:,.1f} juta"
    return f"Rp{v:,.0f}"


def fmt_pct(v):
    return f"{v*100:.1f}%"


def fmt_date(v):
    if isinstance(v, datetime):
        return v.strftime("%d %b %Y")
    try:
        return pd.to_datetime(v).strftime("%d %b %Y")
    except Exception:
        return str(v) if v else "—"


def source_badge(kind: str):
    cls = {
        "Aruna Source": "badge-source",
        "Derived": "badge-derived",
        "Assumption": "badge-assumption",
        "Model Layer": "badge-derived",
    }.get(kind, "badge-source")
    return f'<span class="badge {cls}">{kind}</span>'


def section_header(kicker: str, title: str):
    st.markdown(f'<div class="section-kicker">{kicker}</div><div class="section-title">{title}</div>', unsafe_allow_html=True)


def signal(title: str, body: str, kind=""):
    st.markdown(
        f'<div class="signal {kind}"><div class="signal-title">{title}</div><div class="signal-body">{body}</div></div>',
        unsafe_allow_html=True,
    )


def plot_layout(fig, height=390):
    # Force every Plotly text layer to a dark, readable color.
    # This avoids Streamlit/browser theme overrides making SVG text white.
    DARK = "#17384C"
    GRID = "#8B9298"
    fig.update_layout(
        template="plotly_white",
        height=height,
        paper_bgcolor="rgba(0,0,0,0)",
        plot_bgcolor="rgba(255,255,255,0)",
        margin=dict(l=70, r=70, t=48, b=58),
        font=dict(family="Arial, sans-serif", color=DARK, size=12),
        title_font=dict(family="Arial, sans-serif", size=14, color=DARK),
        legend=dict(
            title=dict(font=dict(color=DARK, size=11)),
            font=dict(color=DARK, size=11),
            bgcolor="rgba(255,255,255,0.75)",
        ),
        coloraxis_colorbar=dict(
            title=dict(font=dict(color=DARK, size=11)),
            tickfont=dict(color=DARK, size=10),
        ),
        uniformtext_minsize=10,
        uniformtext_mode="hide",
    )
    fig.update_xaxes(
        showgrid=True,
        gridcolor=GRID,
        gridwidth=1,
        zeroline=False,
        tickfont=dict(family="Arial, sans-serif", color=DARK, size=11),
        title_font=dict(family="Arial, sans-serif", color=DARK, size=12),
        tickcolor=DARK,
        linecolor=DARK,
        showline=True,
    )
    fig.update_yaxes(
        showgrid=True,
        gridcolor=GRID,
        gridwidth=1,
        zeroline=False,
        tickfont=dict(family="Arial, sans-serif", color=DARK, size=11),
        title_font=dict(family="Arial, sans-serif", color=DARK, size=12),
        tickcolor=DARK,
        linecolor=DARK,
        showline=True,
    )

    # Explicitly set text colors for traces and annotations where supported.
    for trace in fig.data:
        if hasattr(trace, "textfont"):
            trace.textfont = dict(color=DARK, size=11)
        if hasattr(trace, "insidetextfont"):
            trace.insidetextfont = dict(color=DARK, size=11)
        if hasattr(trace, "outsidetextfont"):
            trace.outsidetextfont = dict(color=DARK, size=11)

    if fig.layout.annotations:
        for ann in fig.layout.annotations:
            ann.font = dict(color=DARK, size=11)

    return fig


PLOT_CONFIG = {
    "displayModeBar": False,
    "displaylogo": False,
    "responsive": True,
}


# ============================================================
# EXCEL DATABASE LAYER
# ============================================================
@st.cache_data(show_spinner=False)
def load_model_bytes(raw: bytes) -> Dict:
    wb = load_workbook(BytesIO(raw), data_only=True)
    cost = wb["Cost"]
    price = wb["Price"]
    proc = wb["Processing"]
    rawmat = wb["Raw Material"]
    product = wb["Product"]
    final_econ = wb["FINAL_02_Ekonomi"]
    final_resource = wb["FINAL_03_Resource"]
    final_control = wb["FINAL_01_Kendali"]

    # ---- Aruna raw tables ----
    raw_rows = []
    for r in range(7, rawmat.max_row + 1):
        dt, vessel, code, barcode, po_name = [rawmat.cell(r, c).value for c in [2,3,4,5,6]]
        yft, bigeye, fish_count, size, grade = [rawmat.cell(r, c).value for c in [7,8,9,10,11]]
        if not isinstance(dt, datetime) or not isinstance(yft, (int, float)):
            continue
        raw_rows.append({
            "Date": dt, "Vessel": vessel, "Code": code, "Barcode": barcode, "PO": po_name,
            "YFT kg": float(yft), "Big Eye kg": _num(bigeye), "Fish Count": _num(fish_count),
            "Size": str(size).strip() if size else "Unknown",
            "Grade": str(grade).strip().replace("GRADE ", "Grade ") if grade else "Unknown",
        })
    raw_df = pd.DataFrame(raw_rows)

    proc_rows = []
    for r in range(3, proc.max_row + 1):
        dt, vessel, barcode, total_rm, fish_count, loin_fresh, yield_v = [proc.cell(r, c).value for c in range(1, 8)]
        if not isinstance(dt, datetime) or not isinstance(total_rm, (int, float)):
            continue
        proc_rows.append({
            "Date": dt, "Vessel": vessel, "Barcode": barcode, "Total RM kg": _num(total_rm),
            "Fish Count": _num(fish_count), "Loin Fresh kg": _num(loin_fresh), "Yield": _num(yield_v),
        })
    proc_df = pd.DataFrame(proc_rows)

    price_rows = []
    for r in range(6, 13):
        price_rows.append({
            "Date": price.cell(r, 1).value,
            "SKU": price.cell(r, 2).value,
            "USD/lbs": _num(price.cell(r, 3).value),
            "USD/kg": _num(price.cell(r, 4).value),
        })
    price_df = pd.DataFrame(price_rows)

    cost_rows = []
    for r in range(8, 15):
        cost_rows.append({
            "SKU": cost.cell(r, 2).value,
            "PO kg": _num(cost.cell(r, 5).value),
            "Price USD/lbs": _num(cost.cell(r, 6).value),
            "Revenue": _num(cost.cell(r, 7).value),
            "RM": _num(cost.cell(r, 8).value),
            "Processing": _num(cost.cell(r, 9).value),
            "Handling": _num(cost.cell(r, 10).value),
            "Logistic": _num(cost.cell(r, 11).value),
            "COGS": _num(cost.cell(r, 13).value),
            "GP": _num(cost.cell(r, 14).value),
            "GP %": _num(cost.cell(r, 15).value),
        })
    cost_sku_df = pd.DataFrame(cost_rows)

    excess_rows = []
    for r in range(18, 24):
        excess_rows.append({
            "Category": str(cost.cell(r, 2).value),
            "Qty kg": _num(cost.cell(r, 5).value),
            "Processing": _num(cost.cell(r, 9).value),
            "Handling": _num(cost.cell(r, 10).value),
            "COGS": _num(cost.cell(r, 13).value),
            "GP": _num(cost.cell(r, 14).value),
        })
    excess_df = pd.DataFrame(excess_rows)

    product_rows = []
    current_desc = None
    for r in range(3, product.max_row + 1):
        desc = product.cell(r, 1).value
        if desc:
            current_desc = str(desc).replace("\n", " ")
        if product.cell(r, 2).value is None:
            continue
        product_rows.append({
            "Product": current_desc,
            "Size": product.cell(r, 2).value,
            "Length": product.cell(r, 3).value,
            "Width": product.cell(r, 4).value,
            "Thickness": product.cell(r, 5).value,
            "Weight pcs/gr": product.cell(r, 6).value,
            "Qty pcs/MC": product.cell(r, 7).value,
            "Package kg": product.cell(r, 8).value,
        })
    product_df = pd.DataFrame(product_rows)

    # ---- Core model values ----
    po = np.array([_num(cost.cell(r, 5).value) for r in range(8,15)] + [_num(final_control["B20"].value)], dtype=float)
    stock = np.array([
        _num(cost["D36"].value), _num(cost["D37"].value), _num(cost["D38"].value),
        _num(cost["D39"].value), _num(cost["D40"].value), 0, 0, 0,
    ], dtype=float)
    prices = np.array([_num(price.cell(r, 3).value) for r in range(6,13)] + [_num(final_control["B19"].value, 2)], dtype=float)
    fx = _num(cost["H2"].value, 16481)
    lbskg = _num(final_control["B13"].value, 2.2046226218)

    whole = proc_df["Total RM kg"].sum()
    loin = proc_df["Loin Fresh kg"].sum()
    y = loin / whole if whole else 0.590674080991595

    rho = np.array([_num(final_econ.cell(r, 8).value) for r in range(4,12)], dtype=float)
    energy = np.array([_num(final_econ.cell(r, 12).value) for r in range(4,12)], dtype=float)
    conv = np.array([_num(final_econ.cell(r, 9).value) for r in range(4,12)], dtype=float)

    resource_ledger = {g: 0.0 for g in FRESH_GROUPS}
    if not raw_df.empty:
        agg = raw_df.groupby(["Size","Grade"], as_index=False)["YFT kg"].sum()
        for _, row in agg.iterrows():
            key = f"{row['Size']} {row['Grade']}"
            if key in resource_ledger:
                resource_ledger[key] += float(row["YFT kg"])
    ledger_vec = np.array([resource_ledger[g] for g in FRESH_GROUPS], dtype=float)
    ledger_share = ledger_vec / ledger_vec.sum() if ledger_vec.sum() else np.array([0,0,0,0], dtype=float)

    wgg = _num(cost["D28"].value, 50136)
    wip_stsk = _num(cost["D31"].value, 8643)
    eligibility = np.array(
        [[_num(final_resource.cell(r, c).value) for c in range(2,10)] for r in range(18,23)], dtype=int
    )

    container = _num(final_control["B7"].value, 19000)
    daily_capacity = _num(final_control["B9"].value, 4500)
    horizon = _num(final_control["B10"].value, 30)
    monthly_cap = _num(final_control["B11"].value, 60000)
    historical_intensity = _num(final_control["B15"].value, 0.4636475344)

    current_wip_excess = excess_df.loc[excess_df["Category"].isin(["WIP ST/SK","WIP GM/PK","EXCESS FG"]), "Qty kg"].sum()
    current_gp = _num(cost["G119"].value, 293469606)
    current_fg = _num(cost["E16"].value, 19976)

    # Current shipment metadata
    customer = cost["D2"].value
    po_number = cost["I3"].value
    shipment_date = cost["D4"].value

    # Data catalog / provenance
    raw_start = raw_df["Date"].min() if not raw_df.empty else None
    raw_end = raw_df["Date"].max() if not raw_df.empty else None
    proc_start = proc_df["Date"].min() if not proc_df.empty else None
    proc_end = proc_df["Date"].max() if not proc_df.empty else None
    data_catalog = pd.DataFrame([
        ["Raw Material", "Supply & size-grade composition", len(raw_df), f"{fmt_date(raw_start)} – {fmt_date(raw_end)}", "Aruna Source"],
        ["Processing", "Whole fish → loin yield", len(proc_df), f"{fmt_date(proc_start)} – {fmt_date(proc_end)}", "Aruna Source"],
        ["Price", "SKU selling price", len(price_df), str(price_df.iloc[0]["Date"]), "Aruna Source"],
        ["Cost", "PO, shipment economics, WIP / excess", len(cost_sku_df)+len(excess_df), f"PO {po_number} | ship {fmt_date(shipment_date)}", "Aruna Source"],
        ["Product", "Product / packaging specification", len(product_df), "Case specification", "Aruna Source"],
        ["FINAL model sheets", "rho, energy, eligibility & solver logic", 5, "Optimization layer", "Model Layer"],
    ], columns=["Source Sheet","Decision Use","Records","Coverage","Status"])

    return {
        "raw_df": raw_df, "proc_df": proc_df, "price_df": price_df, "cost_sku_df": cost_sku_df,
        "excess_df": excess_df, "product_df": product_df, "data_catalog": data_catalog,
        "po": po, "stock": stock, "prices": prices, "fx": fx, "lbskg": lbskg,
        "whole_to_loin_yield": y, "rho": rho, "energy": energy, "conversion_cost_fg": conv,
        "ledger_share": ledger_share, "ledger_vec": ledger_vec,
        "wgg": wgg, "wip_stsk": wip_stsk, "eligibility": eligibility,
        "container": container, "daily_capacity": daily_capacity, "horizon": horizon,
        "monthly_cap": monthly_cap, "historical_intensity": historical_intensity,
        "current_wip_excess": current_wip_excess, "current_gp": current_gp, "current_fg": current_fg,
        "customer": customer, "po_number": po_number, "shipment_date": shipment_date,
        "raw_start": raw_start, "raw_end": raw_end, "proc_start": proc_start, "proc_end": proc_end,
    }


def get_model(uploaded_file):
    if uploaded_file is not None:
        raw = uploaded_file.getvalue()
        source_name = uploaded_file.name
    else:
        raw = DEFAULT_MODEL.read_bytes()
        source_name = DEFAULT_MODEL.name
    return load_model_bytes(raw), source_name


# ============================================================
# OPTIMIZATION ENGINE
# ============================================================
@dataclass
class Solution:
    success: bool
    objective_name: str
    allocation: np.ndarray
    new_fg: np.ndarray
    stock_used: np.ndarray
    shipment_fg: np.ndarray
    resource_available_loin: np.ndarray
    residual_by_resource: np.ndarray
    routing_value: float
    energy_total: float
    energy_intensity: float
    shipment_total: float
    container_fill: float
    profit_max_reference: float
    retention: float


def prepare_problem(model: Dict, cfg: Dict):
    # User dapat mengetik kondisi operasional secara langsung.
    po = np.array(cfg.get("po_input", model["po"]), dtype=float)
    stock = np.array(cfg.get("stock_input", model["stock"]), dtype=float)

    stock_used = np.minimum(stock, po) if cfg["use_stock"] else np.zeros_like(stock)
    net_po = np.maximum(po - stock_used, 0)

    wgg_input = float(cfg.get("wgg_input", model["wgg"]))
    wip_input = float(cfg.get("wip_input", model["wip_stsk"]))

    # Proporsi size-grade mengikuti ledger Aruna; volume shipment dapat diubah user.
    fresh_whole = wgg_input * model["ledger_share"]
    fresh_loin = fresh_whole * model["whole_to_loin_yield"]
    avail = np.r_[fresh_loin, wip_input]

    revenue_fg = model["prices"] * model["lbskg"] * model["fx"]
    contribution_fg = revenue_fg - model["conversion_cost_fg"]
    contribution_loin = contribution_fg * model["rho"]

    contract_min = np.clip(np.array(cfg["contract_min"], dtype=float), 0, po)
    min_new = np.maximum(contract_min - stock_used, 0)
    intensity_cap = model["historical_intensity"] * (1 - cfg["energy_target"])

    return {
        "po": po, "stock_used": stock_used, "net_po": net_po, "avail_loin": avail,
        "contribution_loin": contribution_loin, "min_new_fg": min_new,
        "intensity_cap": intensity_cap,
    }


def build_lp(model: Dict, p: Dict, cfg: Dict):
    n_r, n_s = len(RESOURCE_GROUPS), len(SKUS)
    n = n_r * n_s
    rho, energy, eligibility = model["rho"], model["energy"], model["eligibility"]
    A, b = [], []

    # supply
    for i in range(n_r):
        row = np.zeros(n); row[i*n_s:(i+1)*n_s] = 1
        A.append(row); b.append(p["avail_loin"][i])

    # demand ceiling
    for j in range(n_s):
        row = np.zeros(n)
        for i in range(n_r): row[i*n_s+j] = rho[j]
        A.append(row); b.append(p["net_po"][j])

    # contractual minimum
    for j in range(n_s):
        if p["min_new_fg"][j] > 0:
            row = np.zeros(n)
            for i in range(n_r): row[i*n_s+j] = -rho[j]
            A.append(row); b.append(-p["min_new_fg"][j])

    # container max and optional minimum
    fg_row = np.zeros(n)
    for i in range(n_r):
        for j in range(n_s): fg_row[i*n_s+j] = rho[j]
    A.append(fg_row.copy()); b.append(max(model["container"] - p["stock_used"].sum(), 0))
    required_new = max(cfg["min_fill"] - p["stock_used"].sum(), 0)
    if required_new > 0:
        A.append(-fg_row.copy()); b.append(-required_new)

    # monthly cap
    monthly_limit = model["monthly_cap"] * (model["horizon"] / 26)
    A.append(fg_row.copy()); b.append(max(monthly_limit - p["stock_used"].sum(), 0))

    # optional energy intensity constraint
    if cfg["energy_active"]:
        row = np.zeros(n)
        for i in range(n_r):
            for j in range(n_s):
                row[i*n_s+j] = rho[j] * (energy[j] - p["intensity_cap"])
        A.append(row); b.append(0.0)

    bounds = []
    for i in range(n_r):
        for j in range(n_s):
            bounds.append((0, None if eligibility[i,j] == 1 else 0))

    return np.array(A), np.array(b), bounds


def unpack(model, p, x, name, profit_ref, retention):
    n_r, n_s = len(RESOURCE_GROUPS), len(SKUS)
    alloc = x.reshape(n_r, n_s)
    new_fg = alloc.sum(axis=0) * model["rho"]
    shipment = new_fg + p["stock_used"]
    residual = p["avail_loin"] - alloc.sum(axis=1)
    value = float(np.sum(alloc * p["contribution_loin"][None,:]))
    energy_total = float(np.dot(new_fg, model["energy"]))
    intensity = energy_total / new_fg.sum() if new_fg.sum() else 0
    total = float(shipment.sum())
    return Solution(True, name, alloc, new_fg, p["stock_used"], shipment, p["avail_loin"], residual,
                    value, energy_total, intensity, total, total/model["container"] if model["container"] else 0,
                    profit_ref, retention)


def solve_profit_max(model, p, cfg):
    A, b, bounds = build_lp(model, p, cfg)
    c = np.tile(-p["contribution_loin"], len(RESOURCE_GROUPS))
    res = linprog(c, A_ub=A, b_ub=b, bounds=bounds, method="highs")
    if not res.success: return None, res.message
    val = -float(res.fun)
    return unpack(model, p, res.x, "Profit Max", val, 1.0), None


def solve_balanced(model, p, cfg):
    pm, err = solve_profit_max(model, p, cfg)
    if pm is None: return None, err
    A, b, bounds = build_lp(model, p, cfg)
    profit_coef = np.tile(p["contribution_loin"], len(RESOURCE_GROUPS))
    A = np.vstack([A, -profit_coef]); b = np.r_[b, -cfg["retention"]*pm.routing_value]
    c = -np.ones(len(RESOURCE_GROUPS)*len(SKUS))
    res = linprog(c, A_ub=A, b_ub=b, bounds=bounds, method="highs")
    if not res.success: return None, res.message
    return unpack(model, p, res.x, "Balanced — Excess Min", pm.routing_value, cfg["retention"]), None


def solve_energy(model, p, cfg):
    pm, err = solve_profit_max(model, p, cfg)
    if pm is None: return None, err
    bal, err = solve_balanced(model, p, cfg)
    if bal is None: return None, err
    A, b, bounds = build_lp(model, p, cfg)
    profit_coef = np.tile(p["contribution_loin"], len(RESOURCE_GROUPS))
    A = np.vstack([A, -profit_coef]); b = np.r_[b, -cfg["retention"]*pm.routing_value]
    allocated_min = p["avail_loin"].sum() - bal.residual_by_resource.sum()
    A = np.vstack([A, -np.ones(len(RESOURCE_GROUPS)*len(SKUS))]); b = np.r_[b, -allocated_min]
    energy_coef = np.zeros(len(RESOURCE_GROUPS)*len(SKUS))
    for i in range(len(RESOURCE_GROUPS)):
        for j in range(len(SKUS)):
            energy_coef[i*len(SKUS)+j] = model["rho"][j] * model["energy"][j]
    res = linprog(energy_coef, A_ub=A, b_ub=b, bounds=bounds, method="highs")
    if not res.success: return None, res.message
    return unpack(model, p, res.x, "Energy Efficient", pm.routing_value, cfg["retention"]), None


def solve(model, p, cfg):
    if cfg["objective"] == "Profit Max": return solve_profit_max(model, p, cfg)
    if cfg["objective"] == "Balanced — Excess Min": return solve_balanced(model, p, cfg)
    return solve_energy(model, p, cfg)


# ============================================================
# REPORTING HELPERS
# ============================================================
def plan_df(model, p, sol):
    fulfillment = np.divide(sol.shipment_fg, p["po"], out=np.zeros_like(p["po"]), where=p["po"]>0)
    return pd.DataFrame({
        "SKU": SKUS,
        "PO (kg)": p["po"],
        "Stok FG Terpakai (kg)": sol.stock_used,
        "Rekomendasi Produksi Baru (kg)": sol.new_fg,
        "Shipment (kg)": sol.shipment_fg,
        "Pemenuhan PO": fulfillment,
        "Energi (kWh)": sol.new_fg * model["energy"],
        "Intensitas Energi (kWh/kg)": model["energy"],
        "Nilai Routing/kg Loin": p["contribution_loin"],
    })


def allocation_df(sol):
    return pd.DataFrame(sol.allocation, index=RESOURCE_GROUPS, columns=SKUS)


def recommendation_excel(model, p, sol, cfg):
    out = BytesIO()
    with pd.ExcelWriter(out, engine="openpyxl") as writer:
        plan_df(model, p, sol).to_excel(writer, sheet_name="Production Plan", index=False)
        allocation_df(sol).to_excel(writer, sheet_name="Resource Allocation")
        pd.DataFrame({
            "Metric": ["Objective","Routing Value","Ending Residual Loin kg","Energy Total kWh","Energy Intensity kWh/kg","Shipment kg","Container Fill","Profit Retention"],
            "Value": [sol.objective_name, sol.routing_value, sol.residual_by_resource.sum(), sol.energy_total, sol.energy_intensity, sol.shipment_total, sol.container_fill, sol.retention],
        }).to_excel(writer, sheet_name="Decision Summary", index=False)
        pd.DataFrame([cfg]).to_excel(writer, sheet_name="Scenario Inputs", index=False)
    return out.getvalue()


def sankey(sol):
    labels = RESOURCE_GROUPS + SKUS
    source, target, value = [], [], []
    for i in range(len(RESOURCE_GROUPS)):
        for j in range(len(SKUS)):
            if sol.allocation[i,j] > 1:
                source.append(i); target.append(len(RESOURCE_GROUPS)+j); value.append(float(sol.allocation[i,j]))
    fig = go.Figure(go.Sankey(
        arrangement="snap",
        node=dict(label=labels, pad=20, thickness=17, color=[TEAL_2]*len(RESOURCE_GROUPS)+[NAVY_2]*len(SKUS)),
        link=dict(source=source, target=target, value=value, color="rgba(22,151,166,.20)"),
    ))
    fig.update_layout(height=500, margin=dict(l=10,r=10,t=15,b=10), paper_bgcolor="rgba(0,0,0,0)", font=dict(color=INK,size=11))
    return fig



# ============================================================
# V5 UI — RAW MATERIAL ALLOCATION CONTROL TOWER
# ============================================================
LOGO_PATH = Path(__file__).with_name("aruna_logo.png")

NAVY = "#075F92"
NAVY_2 = "#0B3E5B"
TEAL = "#0D8EAF"
TEAL_2 = "#52B8CB"
AMBER = "#FF6B1A"
INK = "#17384C"
MUTED = "#71828E"
BG = "#F7F6F2"
BORDER = "#E4E7E9"
SOFT_AMBER = "#FFF1E7"

st.markdown(
    f"""
    <style>
    .stApp {{background:{BG};}}
    .block-container {{max-width:1520px; padding-top:1.0rem; padding-bottom:3rem;}}
    [data-testid="stSidebar"] {{background:linear-gradient(180deg,#075F92 0%,#064B73 100%)!important; border-right:0!important; min-width:270px!important; max-width:270px!important;}}
    [data-testid="stSidebar"] label,
[data-testid="stSidebar"] p,
[data-testid="stSidebar"] .stMarkdown,
[data-testid="stSidebar"] .stCaption {{
    color:#F6FBFD !important;
}}
[data-testid="stSidebar"] input,
[data-testid="stSidebar"] textarea,
[data-testid="stSidebar"] [data-baseweb="select"] * {{
    color:#17384C !important;
}}
[data-testid="stDataFrame"] *,
[data-testid="stDataEditor"] * {{
    color:#17384C;
}}
.stSelectbox label,
.stRadio label,
.stSlider label,
.stNumberInput label,
.stToggle label,
.stFileUploader label {{
    color:#17384C !important;
}}
    [data-testid="stSidebar"] .stCaption {{color:#CDE6F1!important;}}
    [data-testid="stSidebar"] [data-testid="stFileUploader"] {{background:rgba(255,255,255,.08); border:1px solid rgba(255,255,255,.16); border-radius:13px; padding:.4rem;}}
    .aruna-logo-box {{background:white;border-radius:14px;padding:.45rem .65rem;margin:.05rem 0 .75rem;}}
    .side-title {{font-size:1.08rem;font-weight:850;color:white;letter-spacing:-.02em;}}
    .side-sub {{font-size:.75rem;color:#D6EAF3;line-height:1.42;margin:.2rem 0 .7rem;}}
    .scenario-box {{background:rgba(255,255,255,.10);border:1px solid rgba(255,255,255,.18);border-radius:14px;padding:.7rem .78rem;margin:.7rem 0;}}
    .scenario-box b {{color:white;font-size:.82rem;}} .scenario-box div {{font-size:.7rem;color:#D8EBF3;line-height:1.5;margin-top:.2rem;}}
    .topbar {{background:white;border:1px solid {BORDER};border-bottom:4px solid {AMBER};border-radius:18px;padding:.95rem 1.15rem;box-shadow:0 4px 16px rgba(20,61,82,.05);margin-bottom:1rem;}}
    .top-kicker {{font-size:.7rem;color:{TEAL};font-weight:850;text-transform:uppercase;letter-spacing:.11em;}}
    .top-title {{font-size:1.72rem;color:{NAVY_2};font-weight:880;letter-spacing:-.035em;line-height:1.12;margin:.15rem 0;}}
    .top-sub {{font-size:.9rem;color:{MUTED};line-height:1.48;max-width:1050px;}}
    .pill {{display:inline-block;margin:.48rem .28rem 0 0;padding:.34rem .62rem;border-radius:999px;font-size:.7rem;font-weight:780;background:#EAF6F9;color:#08728B;}}
    .pill.orange {{background:#FFF0E6;color:#D05814;}}
    .kpi {{background:white;border:1px solid {BORDER};border-top:4px solid {TEAL};border-radius:15px;padding:.82rem .9rem;min-height:112px;box-shadow:0 3px 10px rgba(20,61,82,.035);}}
    .kpi.orange {{border-top-color:{AMBER};}} .kpi-label {{font-size:.68rem;color:{MUTED};font-weight:820;text-transform:uppercase;letter-spacing:.055em;}}
    .kpi-value {{font-size:1.55rem;color:{NAVY_2};font-weight:880;letter-spacing:-.03em;line-height:1.1;margin:.22rem 0 .12rem;}}
    .kpi-value.text {{font-size:1.12rem;line-height:1.2;}} .kpi-note {{font-size:.7rem;color:#647C87;line-height:1.35;}}
    .dashboard-panel {{background:white;border:1px solid {BORDER};border-radius:16px;padding:.85rem .9rem;box-shadow:0 3px 10px rgba(20,61,82,.03);}}
    .panel-head {{font-size:.94rem;font-weight:830;color:{NAVY_2};}} .panel-sub2 {{font-size:.75rem;color:{MUTED};line-height:1.4;margin:.15rem 0 .5rem;}}
    .scenario-card-v5 {{background:white;border:1px solid {BORDER};border-radius:15px;padding:.82rem .88rem;min-height:145px;}}
    .scenario-card-v5.active {{border:2px solid {AMBER};background:#FFF9F4;}}
    .sc-title {{font-size:.92rem;font-weight:850;color:{NAVY_2};}} .sc-value {{font-size:1.28rem;font-weight:880;color:{NAVY_2};margin:.3rem 0;}}
    .sc-note {{font-size:.72rem;color:{MUTED};line-height:1.45;}}
    [data-testid="stForm"] {{background:white!important;border:1px solid {BORDER}!important;border-radius:17px!important;padding:.9rem 1rem!important;}}
    [data-testid="stFormSubmitButton"] button {{background:{AMBER}!important;color:white!important;border:0!important;border-radius:12px!important;min-height:48px!important;font-weight:880!important;}}
    .stButton > button {{background:white!important;color:{NAVY_2}!important;border:1px solid #BDD4DE!important;border-radius:11px!important;min-height:40px!important;font-weight:770!important;}}
    .stButton > button[kind="primary"] {{background:{AMBER}!important;color:white!important;border-color:{AMBER}!important;}}
    .stDownloadButton > button {{border-radius:11px!important;min-height:40px!important;font-weight:770!important;}}
    div[data-testid="stDataFrame"] {{border-radius:13px;overflow:hidden;border:1px solid {BORDER};}}
    .section-title {{font-size:1.12rem!important;}} .section-kicker {{font-weight:850!important;}}
    </style>
    """,
    unsafe_allow_html=True,
)

# ---- Final readability overrides ----
# Keep text dark on light content surfaces and white only on the blue sidebar.
st.markdown(
    """
    <style>
    .dashboard-panel,
    .kpi,
    .scenario-card-v5,
    .topbar,
    [data-testid="stForm"] {
        color: #17384C !important;
    }
    .dashboard-panel *,
    .kpi *,
    .scenario-card-v5 *,
    .topbar *,
    [data-testid="stForm"] label,
    [data-testid="stForm"] p,
    [data-testid="stForm"] span {
        color: #17384C;
    }
    [data-testid="stSidebar"] {
        color: #F6FBFD !important;
    }
    [data-testid="stSidebar"] label,
    [data-testid="stSidebar"] p,
    [data-testid="stSidebar"] .stMarkdown,
    [data-testid="stSidebar"] .stCaption {
        color: #F6FBFD !important;
    }
    [data-testid="stSidebar"] input {
        color: #17384C !important;
        background: #FFFFFF !important;
    }

    /* Streamlit widgets on light backgrounds */
    .stAlert, [data-testid="stAlert"] {
        color: #17384C !important;
    }
    .stAlert p, [data-testid="stAlert"] p {
        color: #17384C !important;
    }
    </style>
    """,
    unsafe_allow_html=True,
)


def label_objective(name: str) -> str:
    return {"Profit Max":"Nilai Ekonomi Maksimum","Balanced — Excess Min":"Seimbang: Kurangi Excess","Energy Efficient":"Efisiensi Energi"}.get(name,name)


def kpi(label,value,note="",orange=False,text=False):
    cls="kpi orange" if orange else "kpi"; vcls="kpi-value text" if text else "kpi-value"
    st.markdown(f'<div class="{cls}"><div class="kpi-label">{label}</div><div class="{vcls}">{value}</div><div class="kpi-note">{note}</div></div>',unsafe_allow_html=True)


def explain(model,p,sol):
    cards=[]; econ=p["contribution_loin"]; j=int(np.argmax(econ))
    cards.append((f"Prioritaskan {SKUS[j]}",f"Nilai routing paling tinggi pada kondisi aktif: {fmt_rp(econ[j])}/kg loin."))
    coverage=np.divide(sol.stock_used,p["po"],out=np.zeros_like(p["po"]),where=p["po"]>0)
    if coverage.max()>0:
        k=int(np.argmax(coverage)); cards.append((f"Gunakan stok {SKUS[k]} lebih dulu",f"Stok sudah menutup sekitar {fmt_pct(coverage[k])} dari PO sehingga produksi ulang dapat dikurangi."))
    if sol.residual_by_resource.sum()>1:
        i=int(np.argmax(sol.residual_by_resource)); cards.append(("Tindak lanjuti sisa resource",f"{fmt_kg(sol.residual_by_resource[i])} dari {RESOURCE_GROUPS[i]} belum memperoleh rute pada shipment ini."))
    return cards


def goto(name): st.session_state["_pending_nav"]=name

def set_obj(name):
    c=st.session_state["scenario_cfg"].copy();c["objective"]=name;st.session_state["scenario_cfg"]=c

def reset_data():
    c=st.session_state["scenario_cfg"].copy();c.update({"objective":"Balanced — Excess Min","use_stock":True,"retention":.95,"min_fill":0.0,"energy_target":0.0,"energy_active":False,"contract_min":[0.0]*len(SKUS),"wgg_input":float(model["wgg"]),"wip_input":float(model["wip_stsk"]),"po_input":model["po"].astype(float).tolist(),"stock_input":model["stock"].astype(float).tolist()});st.session_state["scenario_cfg"]=c

def preset(kind):
    if kind=="reset":reset_data();return
    c=st.session_state["scenario_cfg"].copy()
    if kind=="supply80":c["wgg_input"]=float(model["wgg"])*.8;c["wip_input"]=float(model["wip_stsk"])*.8
    elif kind=="po120":c["po_input"]=(model["po"].astype(float)*1.2).tolist()
    elif kind=="full":c["min_fill"]=float(model["container"])
    elif kind=="energy10":c["energy_active"]=True;c["energy_target"]=.10
    st.session_state["scenario_cfg"]=c


with st.sidebar:
    st.markdown('<div class="aruna-logo-box">',unsafe_allow_html=True);st.image(str(LOGO_PATH),width=132);st.markdown('</div>',unsafe_allow_html=True)
    st.markdown('<div class="side-title">Allocation Control Tower</div>',unsafe_allow_html=True)
    st.markdown('<div class="side-sub">Strategi alokasi raw material → SKU</div>',unsafe_allow_html=True)
    with st.expander("Database",expanded=False): uploaded=st.file_uploader("Upload Excel Aruna terbaru",type=["xlsx"])

model,source_name=get_model(uploaded)
PAGES=["Overview","Alokasi SKU","Simulasi","Dampak","Data & Asumsi"]
source_key=f"{source_name}|{model['po_number']}|{model['wgg']}|{model['wip_stsk']}"
if st.session_state.get("_source_key")!=source_key:
    st.session_state["_source_key"]=source_key
    st.session_state["scenario_cfg"]={"objective":"Balanced — Excess Min","use_stock":True,"retention":.95,"min_fill":0.0,"energy_target":0.0,"energy_active":False,"contract_min":[0.0]*len(SKUS),"wgg_input":float(model["wgg"]),"wip_input":float(model["wip_stsk"]),"po_input":model["po"].astype(float).tolist(),"stock_input":model["stock"].astype(float).tolist()}
    st.session_state["decision_log"]=[]
cfg=st.session_state["scenario_cfg"].copy()
if "nav_page" not in st.session_state:st.session_state["nav_page"]="Overview"
if st.session_state.get("_pending_nav") in PAGES:st.session_state["nav_page"]=st.session_state.pop("_pending_nav")
with st.sidebar:
    page=st.radio("Menu",PAGES,key="nav_page")
    st.markdown(f'<div class="scenario-box"><b>Skenario Aktif</b><div>{label_objective(cfg["objective"])}<br>Retention {cfg["retention"]:.0%}<br>PO #{model["po_number"]}</div></div>',unsafe_allow_html=True)
    st.caption("Input utama berada di menu Simulasi.")

p=prepare_problem(model,cfg);sol,err=solve(model,p,cfg)
if err:st.error(f"Skenario tidak feasible: {err}");st.stop()

h1,h2=st.columns([1.1,8.9],vertical_alignment="center")
with h1:st.image(str(LOGO_PATH),width=138)
with h2:
    st.markdown(f'''<div class="topbar"><div class="top-kicker">ARUNA — PRODUCTION DECISION SUPPORT SYSTEM</div><div class="top-title">Raw Material Allocation Control Tower</div><div class="top-sub">Menentukan bahan baku mana yang dialokasikan ke SKU mana berdasarkan resource aktual, PO, stok, yield, nilai ekonomi, excess, dan energi.</div><span class="pill">Database Excel Terhubung</span><span class="pill">PO #{model['po_number']}</span><span class="pill orange">Shipment {fmt_date(model['shipment_date'])}</span></div>''',unsafe_allow_html=True)

if page=="Overview":
    section_header("Ringkasan","Kondisi utama sebelum strategi alokasi dipilih")
    cols=st.columns(5);data=[("Whole Fish",fmt_kg(cfg["wgg_input"]),"Resource shipment",False),("WIP ST/SK",fmt_kg(cfg["wip_input"]),"Carry-in",False),("PO Buyer",fmt_kg(sum(cfg["po_input"])),f"{sum(np.array(cfg['po_input'])>0)} SKU aktif",False),("Stok FG",fmt_kg(sum(cfg["stock_input"])),"Dipakai sebelum produksi ulang",False),("WIP + Excess",fmt_kg(model["current_wip_excess"]),"Proxy modal tertahan",True)]
    for col,(lab,val,note,org) in zip(cols,data):
        with col:kpi(lab,val,note,org)
    left,mid,right=st.columns([1.05,1.2,.85])
    with left:
        st.markdown('<div class="dashboard-panel"><div class="panel-head">Komposisi Raw Material</div><div class="panel-sub2">Size–grade supply dari ledger Aruna</div>',unsafe_allow_html=True)
        mix=model["raw_df"].groupby(["Size","Grade"],as_index=False)["YFT kg"].sum();mix["Resource"]=mix["Size"]+" • "+mix["Grade"]
        fig=px.treemap(mix,path=["Resource"],values="YFT kg",color="YFT kg",color_continuous_scale=["#DFF2F7",TEAL]);fig.update_coloraxes(showscale=False);st.plotly_chart(plot_layout(fig,335),use_container_width=True,config=PLOT_CONFIG);st.markdown('</div>',unsafe_allow_html=True)
    with mid:
        st.markdown('<div class="dashboard-panel"><div class="panel-head">PO vs Existing Stock</div><div class="panel-sub2">Kebutuhan yang benar-benar perlu diproduksi</div>',unsafe_allow_html=True)
        dem=pd.DataFrame({"SKU":SKUS,"PO":p["po"],"Stok FG":np.array(cfg["stock_input"],dtype=float)});dem=dem[dem["PO"]>0].melt("SKU",value_vars=["PO","Stok FG"],var_name="Seri",value_name="kg")
        fig=px.bar(dem,x="kg",y="SKU",color="Seri",orientation="h",barmode="group",color_discrete_map={"PO":NAVY_2,"Stok FG":AMBER});fig.update_layout(xaxis_title="kg FG",yaxis_title="");st.plotly_chart(plot_layout(fig,335),use_container_width=True,config=PLOT_CONFIG);st.markdown('</div>',unsafe_allow_html=True)
    with right:
        st.markdown('<div class="dashboard-panel"><div class="panel-head">Decision Signal</div><div class="panel-sub2">Apa yang perlu ditindaklanjuti?</div>',unsafe_allow_html=True)
        for t,b in explain(model,p,sol):signal(t,b,"warn" if "sisa" in t.lower() else "good")
        st.markdown('</div>',unsafe_allow_html=True)
    section_header("Skenario Aktif","Ringkasan rekomendasi saat ini")
    a,b,c,d=st.columns(4)
    with a:kpi("Prioritas",label_objective(sol.objective_name),"Ubah di menu Simulasi",True,True)
    with b:kpi("Shipment",fmt_kg(sol.shipment_total),f"{fmt_pct(sol.container_fill)} kapasitas")
    with c:kpi("Sisa Resource",fmt_kg(sol.residual_by_resource.sum()),"Belum memperoleh rute")
    with d:kpi("Intensitas Energi",f"{sol.energy_intensity:.3f} kWh/kg",f"Historis {model['historical_intensity']:.3f}")
    x,y,z=st.columns([1.2,1,1])
    with x:st.button("ATUR INPUT & HITUNG →",on_click=goto,args=("Simulasi",),type="primary",use_container_width=True)
    with y:st.button("Lihat Strategi Alokasi",on_click=goto,args=("Alokasi SKU",),use_container_width=True)
    with z:st.button("Bandingkan Dampak",on_click=goto,args=("Dampak",),use_container_width=True)

elif page=="Alokasi SKU":
    section_header("Strategi Alokasi","Raw material mana diarahkan ke SKU mana?")
    a,b,c,d=st.columns(4)
    with a:kpi("Produksi Baru",fmt_kg(sol.new_fg.sum()),"Hasil optimizer")
    with b:kpi("Stok FG Dipakai",fmt_kg(sol.stock_used.sum()),"Mengurangi produksi ulang")
    with c:kpi("Shipment",fmt_kg(sol.shipment_total),f"{fmt_pct(sol.container_fill)} kapasitas")
    with d:kpi("Residual",fmt_kg(sol.residual_by_resource.sum()),"Potential WIP / route lain",True)
    l,r=st.columns([1.4,.8])
    with l:
        st.markdown('<div class="dashboard-panel"><div class="panel-head">Peta Resource → SKU</div><div class="panel-sub2">Semakin tebal aliran, semakin besar kg loin yang dialokasikan.</div>',unsafe_allow_html=True);st.plotly_chart(sankey(sol),use_container_width=True,config=PLOT_CONFIG);st.markdown('</div>',unsafe_allow_html=True)
    with r:
        st.markdown('<div class="dashboard-panel"><div class="panel-head">Aksi yang Disarankan</div><div class="panel-sub2">Alasan rekomendasi agar model tidak menjadi black box.</div>',unsafe_allow_html=True)
        for t,b in explain(model,p,sol):signal(t,b,"warn" if "sisa" in t.lower() else "good")
        st.markdown('</div>',unsafe_allow_html=True)
    section_header("Rencana Produksi SKU","Berapa kg yang sebaiknya dibuat?")
    full=plan_df(model,p,sol);simple=full[["SKU","PO (kg)","Stok FG Terpakai (kg)","Rekomendasi Produksi Baru (kg)","Shipment (kg)","Pemenuhan PO"]]
    st.dataframe(simple.style.format({"PO (kg)":"{:,.1f}","Stok FG Terpakai (kg)":"{:,.1f}","Rekomendasi Produksi Baru (kg)":"{:,.1f}","Shipment (kg)":"{:,.1f}","Pemenuhan PO":"{:.1%}"}),hide_index=True,use_container_width=True)
    with st.expander("Lihat matriks alokasi kg loin"):st.dataframe(allocation_df(sol).style.format("{:,.1f}"),use_container_width=True)
    aa,bb,cc=st.columns([1.15,1,1])
    with aa:
        if st.button("SETUJUI RENCANA ALOKASI",type="primary",use_container_width=True):
            st.session_state.setdefault("decision_log",[]).append({"Waktu":datetime.now().strftime("%d-%m-%Y %H:%M:%S"),"Prioritas":label_objective(sol.objective_name),"Shipment (kg)":round(sol.shipment_total,2),"Sisa Resource (kg)":round(sol.residual_by_resource.sum(),2)});st.success("Rencana alokasi masuk ke Decision Log.")
    with bb:st.button("Ubah Skenario",on_click=goto,args=("Simulasi",),use_container_width=True)
    with cc:st.download_button("Download Rencana Excel",recommendation_excel(model,p,sol,cfg),"Aruna_DSS_Allocation_Plan.xlsx","application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",use_container_width=True)

elif page=="Simulasi":
    section_header("Simulation & Decision Input","Ubah kondisi aktual lalu hitung strategi alokasi")
    st.info("✏️ Halaman ini adalah area INPUT. Whole Fish, WIP, PO, stok, minimum kontrak, prioritas, shipment, dan target energi dapat diubah.")
    q=st.columns(5)
    for col,label,kind in zip(q,["Reset","Supply -20%","PO +20%","Full Container","Energi -10%"],["reset","supply80","po120","full","energy10"]):
        with col:st.button(label,on_click=preset,args=(kind,),use_container_width=True)
    l,r=st.columns([1.45,.75])
    with l:
        with st.form("sim_form"):
            st.markdown("### 1. Resource aktual");c1,c2=st.columns(2)
            with c1:wgg=st.number_input("Whole Fish tersedia (kg)",min_value=0.0,value=float(cfg["wgg_input"]),step=500.0)
            with c2:wip=st.number_input("WIP ST/SK tersedia (kg loin)",min_value=0.0,value=float(cfg["wip_input"]),step=250.0)
            use_stock=st.toggle("Gunakan stok FG terlebih dahulu",value=cfg["use_stock"])
            st.markdown("### 2. PO & stok per SKU")
            inp=pd.DataFrame({"SKU":SKUS,"PO Buyer (kg)":np.array(cfg["po_input"],dtype=float),"Stok FG (kg)":np.array(cfg["stock_input"],dtype=float),"Minimum Kontrak (kg)":np.array(cfg["contract_min"],dtype=float)})
            edited=st.data_editor(inp,hide_index=True,use_container_width=True,disabled=["SKU"],column_config={"PO Buyer (kg)":st.column_config.NumberColumn(min_value=0.0,step=50.0),"Stok FG (kg)":st.column_config.NumberColumn(min_value=0.0,step=50.0),"Minimum Kontrak (kg)":st.column_config.NumberColumn(min_value=0.0,step=50.0)})
            st.markdown("### 3. Prioritas keputusan");opts=["Profit Max","Balanced — Excess Min","Energy Efficient"]
            obj=st.radio("Prioritas",opts,index=opts.index(cfg["objective"]),format_func=label_objective,horizontal=True)
            retention=st.slider("Nilai ekonomi yang dipertahankan",80,100,int(round(cfg["retention"]*100)),1,format="%d%%")/100
            min_fill=st.slider("Minimum shipment (kg FG)",0,int(model["container"]),int(cfg["min_fill"]),500)
            energy_active=st.toggle("Aktifkan batas intensitas energi",value=cfg["energy_active"])
            energy_target=st.slider("Target penurunan intensitas energi",0,25,int(round(cfg["energy_target"]*100)),1,format="%d%%",disabled=not energy_active)/100
            run=st.form_submit_button("HITUNG STRATEGI ALOKASI",type="primary",use_container_width=True)
        if run:
            cfg.update({"objective":obj,"use_stock":use_stock,"retention":retention,"min_fill":float(min_fill),"energy_active":energy_active,"energy_target":energy_target,"wgg_input":float(wgg),"wip_input":float(wip),"po_input":edited["PO Buyer (kg)"].astype(float).tolist(),"stock_input":edited["Stok FG (kg)"].astype(float).tolist(),"contract_min":edited["Minimum Kontrak (kg)"].astype(float).tolist()});st.session_state["scenario_cfg"]=cfg;st.rerun()
    with r:
        st.markdown('<div class="dashboard-panel"><div class="panel-head">Preview Hasil</div><div class="panel-sub2">Output dari skenario aktif.</div>',unsafe_allow_html=True);kpi("Prioritas",label_objective(sol.objective_name),"Objective aktif",True,True);st.markdown('<br>',unsafe_allow_html=True);kpi("Shipment",fmt_kg(sol.shipment_total),f"{fmt_pct(sol.container_fill)} kapasitas");st.markdown('<br>',unsafe_allow_html=True);kpi("Sisa Resource",fmt_kg(sol.residual_by_resource.sum()),"Potential WIP / excess");st.markdown('</div>',unsafe_allow_html=True)
    st.button("LIHAT HASIL ALOKASI →",on_click=goto,args=("Alokasi SKU",),type="primary",use_container_width=True)

elif page=="Dampak":
    section_header("Decision Impact","Bandingkan strategi sebelum management memilih")
    rows=[]
    for obj in ["Profit Max","Balanced — Excess Min","Energy Efficient"]:
        cc=cfg.copy();cc["objective"]=obj;pp=prepare_problem(model,cc);ss,ee=solve(model,pp,cc)
        if ss:rows.append({"key":obj,"Skenario":label_objective(obj),"Nilai":ss.routing_value,"Sisa":ss.residual_by_resource.sum(),"Intensitas":ss.energy_intensity,"Shipment":ss.shipment_total})
    cols=st.columns(3)
    for col,row in zip(cols,rows):
        active=row["key"]==cfg["objective"]
        with col:
            st.markdown(f'<div class="scenario-card-v5 {"active" if active else ""}"><div class="sc-title">{row["Skenario"]}</div><div class="sc-value">{fmt_rp(row["Nilai"])}</div><div class="sc-note">Sisa: <b>{fmt_kg(row["Sisa"])}</b><br>Intensitas: <b>{row["Intensitas"]:.3f} kWh/kg</b><br>Shipment: <b>{fmt_kg(row["Shipment"])}</b></div></div>',unsafe_allow_html=True)
            st.button("Skenario Aktif" if active else "Gunakan Skenario",disabled=active,on_click=set_obj,args=(row["key"],),key=f'sc_{row["key"]}',use_container_width=True)
    frontier=[]
    for ret in [1,.99,.98,.97,.95,.92,.90]:
        cc=cfg.copy();cc["objective"]="Balanced — Excess Min";cc["retention"]=ret;pp=prepare_problem(model,cc);ss,ee=solve(model,pp,cc)
        if ss:frontier.append({"Retention":ret,"Sisa Resource":ss.residual_by_resource.sum(),"Nilai":ss.routing_value})
    fdf=pd.DataFrame(frontier);l,r=st.columns(2)
    with l:
        fig=px.line(fdf,x="Retention",y="Sisa Resource",markers=True,color_discrete_sequence=[AMBER]);fig.update_xaxes(tickformat=".0%");fig.update_layout(xaxis_title="Nilai ekonomi dipertahankan",yaxis_title="Sisa loin (kg)");st.plotly_chart(plot_layout(fig,380),use_container_width=True,config=PLOT_CONFIG)
    with r:
        sdf=pd.DataFrame(rows);fig=px.scatter(sdf,x="Intensitas",y="Nilai",size="Shipment",text="Skenario",color="Sisa",color_continuous_scale=["#DFF2F7",AMBER]);fig.update_traces(textposition="top center");st.plotly_chart(plot_layout(fig,380),use_container_width=True,config=PLOT_CONFIG)
    st.markdown(
        '<div style="background:#FFF9B8;border:1px solid #E8D95A;border-radius:10px;'
        'padding:0.85rem 1rem;color:#17384C!important;font-size:0.82rem;line-height:1.45;">'
        '<b>Catatan:</b> Nilai Routing adalah decision metric untuk membandingkan '
        'alternatif alokasi, bukan GP akuntansi PC3.'
        '</div>',
        unsafe_allow_html=True,
    )

else:
    section_header("Data & Asumsi","Sumber angka dan parameter yang masih perlu divalidasi")
    st.success(f"Database aktif: {source_name}. Simulasi tidak mengubah workbook sumber.")
    cat=model["data_catalog"].rename(columns={"Source Sheet":"Sheet","Decision Use":"Fungsi","Records":"Jumlah Data","Coverage":"Cakupan","Status":"Status"});st.dataframe(cat,hide_index=True,use_container_width=True)
    a,b,c,d=st.columns(4)
    with a:kpi("Raw Material",f"{len(model['raw_df'])} baris",f"{fmt_date(model['raw_start'])}–{fmt_date(model['raw_end'])}")
    with b:kpi("Processing",f"{len(model['proc_df'])} batch","Sumber yield")
    with c:kpi("Whole→Loin Yield",fmt_pct(model["whole_to_loin_yield"]),"Derived dari batch aktual")
    with d:kpi("SKU Harga",f"{len(model['price_df'])} SKU","Price sheet")
    l,r=st.columns(2)
    with l:
        tmp=model["proc_df"].copy();fig=px.line(tmp,x="Date",y="Yield",markers=True,color_discrete_sequence=[TEAL]);fig.update_yaxes(tickformat=".0%");st.plotly_chart(plot_layout(fig,360),use_container_width=True,config=PLOT_CONFIG)
    with r:
        signal("Recovery loin → SKU","Sebagian masih benchmark/model coefficient dan perlu actual recovery Aruna.","warn");signal("Eligibility kualitas","Mapping Size–Grade ke AAA/AA / buyer specification perlu divalidasi.","warn");signal("Energi per SKU","Masih process-based estimate, belum plant-metered kWh/kg.","warn")
